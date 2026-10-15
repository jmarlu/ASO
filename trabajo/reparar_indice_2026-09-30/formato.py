from pathlib import Path
from copy import deepcopy
import re
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table
from docx.opc.constants import RELATIONSHIP_TYPE as RT
W=Path(__file__).resolve().parent
src=Document(W/'antes.docx');tpl=Document(W.parent.parent/'Programación/F01.PC02 PROPUESTA DIDÁCTICA ed4.docx')
for name in ['Heading 1','Heading 2','Heading 3']:
 a=src.styles[name]._element;b=tpl.styles[name]._element
 for tag in ['pPr','rPr']:
  old=a.find(qn('w:'+tag))
  if old is not None:a.remove(old)
  new=deepcopy(b.find(qn('w:'+tag)))
  for n in list(new.findall(qn('w:numPr'))):new.remove(n)
  a.append(new)
 example=next((p for p in tpl.paragraphs if p.style.name==name),None)
 for p in src.paragraphs:
  if p.style.name!=name:continue
  pp=p._p.get_or_add_pPr()
  for tag in ['spacing','jc','pBdr']:
   for old in list(pp.findall(qn('w:'+tag))):pp.remove(old)
   value=example._p.get_or_add_pPr().find(qn('w:'+tag)) if example is not None else None
   if value is not None:pp.append(deepcopy(value))
  p.paragraph_format.keep_with_next=True
  for r in p.runs:
   for tag in ['color','sz','szCs','b','bCs','rStyle']:
    if r._r.rPr is not None:
     for n in list(r._r.rPr.findall(qn('w:'+tag))):r._r.rPr.remove(n)
# Retain all seven RA rows while restoring the template's grouped header.
old=next(t for t in src.tables if 'Modo: CE compartidos' in '\n'.join(c.text for r in t.rows for c in r.cells))
data=[[c.text for c in row.cells] for row in old.rows[1:]]
model=tpl.tables[14]
elem=deepcopy(model._tbl);old._tbl.addprevious(elem);old._tbl.getparent().remove(old._tbl)
t=Table(elem,src._body)
body_row=deepcopy(t.rows[2]._tr)
for row in list(t.rows)[2:]:t._tbl.remove(row._tr)
for _ in data:t._tbl.append(deepcopy(body_row))
def put(cell,value):
 p=cell.paragraphs[0];fmt=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
 for node in list(p._p):
  if node.tag!=qn('w:pPr'):p._p.remove(node)
 for extra in list(cell._tc.findall(qn('w:p')))[1:]:cell._tc.remove(extra)
 r=p.add_run(value)
 if fmt is not None:r._r.insert(0,fmt)
 r.font.name='Arial';r.font.size=Pt(9)
 p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
put(t.cell(0,2),'Empresa*')
for i,vals in enumerate(data,2):
 for j,value in enumerate(vals):put(t.cell(i,j),value)
for i,row in enumerate(t.rows):
 pr=row._tr.get_or_add_trPr()
 for tag in ['trHeight','tblHeader']:
  for n in list(pr.findall(qn('w:'+tag))):pr.remove(n)
 if i<2:pr.append(OxmlElement('w:tblHeader'))
 pr.append(OxmlElement('w:cantSplit'))
# Consistent decimal separator without changing numeric values.
for node in src._element.iter(qn('w:t')):
 if node.text:node.text=re.sub(r'(\d)\.(\d{2})(\s*%)',r'\1,\2\3',node.text)
# Make the already provided source URLs clickable, preserving their wording.
for p in src.paragraphs:
 for r in list(p.runs):
  text=r.text
  match=re.search(r'https://\S+',text)
  if not match:continue
  url=match.group(0);start,end=match.span();r.text=text[:start]
  link=OxmlElement('w:hyperlink');link.set(qn('r:id'),src.part.relate_to(url,RT.HYPERLINK,is_external=True))
  rr=OxmlElement('w:r');pr=deepcopy(r._r.rPr) if r._r.rPr is not None else OxmlElement('w:rPr')
  for tag in ['color','u']:
   for n in list(pr.findall(qn('w:'+tag))):pr.remove(n)
  color=OxmlElement('w:color');color.set(qn('w:val'),'0000FF');pr.append(color)
  u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u)
  rr.append(pr);tt=OxmlElement('w:t');tt.text=url;rr.append(tt);link.append(rr);r._r.addnext(link)
  if text[end:]:p.add_run(text[end:])
src.save(W/'formato.docx')
print('Formato preparado; 7 filas de RA conservadas.')
