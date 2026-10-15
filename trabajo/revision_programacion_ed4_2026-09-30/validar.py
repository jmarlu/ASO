from pathlib import Path
from decimal import Decimal
from zipfile import ZipFile
import re,json
from lxml import etree
from docx import Document
from docx.oxml.ns import qn
import pymupdf
W=Path(__file__).resolve().parent;R=W.parent.parent
old=Document(W/'original.docx');d=Document(W/'final.docx');tpl=Document(R/'Programación/F01.PC02 PROPUESTA DIDÁCTICA ed4.docx')
checks=[]
def check(name,cond):
 assert cond,name
 checks.append(name)
def txt(t):return '\n'.join(' | '.join(c.text for c in row.cells) for row in t.rows)
check('32 tablas independientes',len(d.tables)==32)
nt=d.tables[0];tt=tpl.tables[0]
check('Normativa: cuatro columnas y doce filas como la plantilla',len(nt.columns)==4 and len(nt.rows)==12)
check('Normativa: mismos encabezados',[c.text for c in nt.rows[0].cells]==[c.text for c in tt.rows[0].cells])
check('Normativa: mismos bloques y agrupaciones',[r.cells[0].text for r in nt.rows]==[r.cells[0].text for r in tt.rows])
check('Normativa: mismas celdas combinadas',len(nt._tbl.xpath('.//w:vMerge'))==len(tt._tbl.xpath('.//w:vMerge')))
check('Normativa: mismas anchuras de columnas',all(abs(a.width-b.width)<1200 for a,b in zip(nt.columns,tt.columns)))
check('Normativa: modalidad presencial','Normativa de la modalidad presencial' in nt.cell(9,2).text)
check('Normativa: referencias completas','500/2024' in txt(nt) and '114/2025' in txt(nt) and 'Formación Profesional.' in nt.cell(5,3).text)
links=[]
for n in nt._tbl.xpath('.//w:hyperlink[@r:id]'):
 links.append(d.part.rels[n.get(qn('r:id'))].target_ref)
oldlinks=[tpl.part.rels[n.get(qn('r:id'))].target_ref for n in tt._tbl.xpath('.//w:hyperlink[@r:id]')]
check('Normativa: enlaces importados con su destino original',links==oldlinks[:4])
uds=[t for t in d.tables if t.cell(0,0).text.startswith('Unidad didáctica')]
check('Siete unidades con CLIL',len(uds)==7 and all(txt(t).count('Actividades CLIL')>=1 for t in uds))
check('Siete unidades con sesiones y aplicación provisional',all('Sesiones / horas aula' in txt(t) and 'Temporalización provisional:' in txt(t) for t in uds))
oldra=next(t for t in old.tables if t.cell(0,0).text=='Resultados de Aprendizaje')
newra=next(t for t in d.tables if t.cell(0,0).text=='Resultados de Aprendizaje')
normalize=lambda s: re.sub(r'\s+',' ',s).strip()
check('Se conservan los siete RA y los 62 CE de la tabla curricular',normalize(txt(oldra))==normalize(txt(newra)))
annex=[t for t in d.tables if t.cell(0,0).text.startswith('Resultado de aprendizaje (0374.')]
weights=[];ceweights={}
for i,t in enumerate(annex,1):
 def val(s):return Decimal(s.replace('%','').replace(',','.').strip())
 weights.append(val(t.cell(1,1).text));ceweights[i]={}
 for row in t.rows[3:]:ceweights[i][row.cells[0].text.strip()[0]]=val(row.cells[1].text)
 check(f'RA{i}: los CE suman 100 %',sum(ceweights[i].values())==100)
check('Pesos RA conservados y total 100 %',weights==list(map(Decimal,['18','8','18','9','5','22','20'])) and sum(weights)==100)
plan=next(t for t in d.tables if 'Modo: CE compartidos' in txt(t))
for i,row in enumerate(plan.rows[1:],1):
 ce=re.findall(r'\b([a-j])\b',row.cells[3].text.split('.',1)[1]);pv=Decimal(row.cells[5].text.replace('%','').replace(',','.').strip())
 cv=Decimal(row.cells[4].text.replace('%','').replace(',','.').strip())
 check(f'Empresa RA{i}: porcentaje trazable a los CE y reparto 100 %',sum(ceweights[i][c] for c in ce)==pv and cv+pv==100)
hrs=next(t for t in d.tables if t.cell(0,0).text=='Bloque')
check('Horas: 102 aula + 31 aplicación = 133',all(sum(int(r.cells[j].text) for r in hrs.rows[1:-1])==int(hrs.rows[-1].cells[j].text) for j in [1,2,3]) and hrs.rows[-1].cells[3].text=='133')
languages=next(t for t in d.tables if t.cell(0,0).text.strip().startswith('Módulo profesional') and len(t.rows)==20)
check('Lenguas vehiculares cumplimentadas',all(r.cells[1].text.strip() for r in languages.rows[1:-1]))
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
with ZipFile(W/'final.docx') as z:
 textparts=[]
 for n in z.namelist():
  if n.startswith(('word/header','word/footer')) and n.endswith('.xml'):textparts.append(''.join(etree.fromstring(z.read(n)).xpath('//w:t/text()',namespaces=ns)))
 check('Encabezados y pies: edición 04 uniforme',not any('edición: 01' in s or 'PROGRAMACIÓN DIDÁCTICA' in s for s in textparts))
 root=etree.fromstring(z.read('word/document.xml'))
 check('Índice automático Word','TOC ' in ''.join(root.xpath('//w:instrText/text()',namespaces=ns)))
pdf=pymupdf.open(W/'final.pdf')
check('PDF sin páginas vacías',all(len(p.get_text().strip())>150 for p in pdf))
check('Índice PDF con 31 enlaces a apartados',len(pdf[1].get_links())==31)
check('Ningún marcador de índice pendiente',not any('[[ASO_TOC]]' in p.get_text() for p in pdf))
bodytext='\n'.join(p.get_text() for p in pdf)
check('Plan empresa provisional y alternativa de centro','PLAN PROVISIONAL Y ADAPTABLE' in bodytext and '100 % centro, 0 % empresa' in bodytext)
W.joinpath('validacion.json').write_text(json.dumps({'pages':len(pdf),'checks':checks},ensure_ascii=False,indent=2))
print(f'{len(checks)} comprobaciones correctas; {len(pdf)} páginas PDF.')
for c in checks:print('OK',c)
