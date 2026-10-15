from pathlib import Path
from copy import deepcopy
import re
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table, _Cell
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
W=Path(__file__).resolve().parent
R=W.parent.parent
d=Document(W/'original.docx'); template=Document(R/'Programación/F01.PC02 PROPUESTA DIDÁCTICA ed4.docx')
body=d._element.body;old=list(body);ts=list(d.tables)
def P(i):return Paragraph(old[i],d._body)
def put(i,text):p=P(i);p.text=text;return p
def rm(e):
 if e.getparent() is not None:e.getparent().remove(e)
def after(e,text,style='Normal'):
 p=d.add_paragraph(text,style);e.addnext(p._p);return p._p
def before(e,text,style='Normal'):
 p=d.add_paragraph(text,style);e.addprevious(p._p);return p._p
def cell(c,text,keep=False):
 if keep:
  pp=deepcopy(c.paragraphs[0]._p.pPr) if c.paragraphs[0]._p.pPr is not None else None
  rp=deepcopy(c.paragraphs[0].runs[0]._r.rPr) if c.paragraphs[0].runs and c.paragraphs[0].runs[0]._r.rPr is not None else None
 c.text=str(text)
 for p in c.paragraphs:
  if keep:
   if pp is not None:p._p.insert(0,deepcopy(pp))
   if rp is not None:p.runs[0]._r.insert(0,deepcopy(rp))
  else:
   p.style='Normal';p.alignment=WD_ALIGN_PARAGRAPH.LEFT
   p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1
   for r in p.runs:r.font.size=Pt(10)
 c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(e,headers,rows,widths):
 t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for j,h in enumerate(headers):cell(t.rows[0].cells[j],h);t.columns[j].width=Cm(widths[j])
 for data in rows:
  cs=t.add_row().cells
  for j,v in enumerate(data):cell(cs[j],v)
 for row in t.rows:
  for j,c in enumerate(row.cells):c.width=Cm(widths[j])
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
 t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 for c in t.rows[0].cells:
  sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7E6E6');c._tc.get_or_add_tcPr().append(sh)
  for p in c.paragraphs:
   for r in p.runs:r.bold=True
 e.addnext(t._tbl);return t

def nonum(p):
 pr=p._p.get_or_add_pPr();n=pr.find(qn('w:numPr'))
 if n is not None:pr.remove(n)
 n=OxmlElement('w:numPr');ni=OxmlElement('w:numId');ni.set(qn('w:val'),'0');n.append(ni);pr.append(n)

put(18,'Curso: 2026/2027 · Adaptación a edición 04: 30/09/2026')
put(16,'Código: 0374    Nº horas: 133    Modalidad: Presencial');put(13,'MÓDULO PROFESIONAL')
for part in d.part.package.parts:
 if hasattr(part,'element') and str(part.partname).startswith(('/word/header','/word/footer')):
  for n in part.element.iter(qn('w:t')):
   if n.text:n.text=n.text.replace('PROGRAMACIÓN DIDÁCTICA','PROPUESTA DIDÁCTICA').replace('Nº edición: 01','Nº edición: 04')
heads={61:'0. PROPUESTAS DE MEJORA DEL CURSO ANTERIOR',63:'1. INTRODUCCIÓN Y JUSTIFICACIÓN TEÓRICA CONTEXTUALIZADA',65:'1.1. Normativa de referencia',69:'1.2. Identificación del título',83:'1.3. Cuadro horario',95:'1.4. Relación entre los estándares de competencia y los módulos del ciclo formativo',100:'1.5. Identificación del módulo / unidad formativa',103:'1.6. Características del alumnado',112:'1.7. Del sector productivo',128:'2. DESARROLLO CURRICULAR',129:'2.1. Competencias y objetivos',133:'2.2. Resultados de aprendizaje y criterios de evaluación',139:'2.3. Relación entre competencias y resultados de aprendizaje',142:'2.4. Contenidos básicos de referencia del módulo',145:'3. METODOLOGÍA DIDÁCTICA',150:'3.1. Estrategias metodológicas',153:'3.2. Líneas pedagógicas del centro',167:'3.3. Recursos',170:'4. UNIDADES',199:'5. PROYECTO INTERMODULAR',200:'5.1. Especificaciones',204:'5.2. Distribución de elementos curriculares',209:'6. PLAN FORMATIVO EN EMPRESA',216:'7. EVALUACIÓN Y CALIFICACIÓN',224:'7.1. Instrumentos de evaluación',227:'7.2. Criterios de calificación',235:'7.3. Recuperación de los resultados de aprendizaje',244:'7.4. Criterios de inasistencia',250:'8. EVALUACIÓN DE LA PRÁCTICA DOCENTE',262:'9. MEDIDAS DE ATENCIÓN A LA DIVERSIDAD',276:'10. ANEXO I'}
for i,title in heads.items():
 p=put(i,title);lev=2 if re.match(r'^\d+\.\d+\.',title) else 1;p.style=f'Heading {lev}';nonum(p)
 p.paragraph_format.keep_with_next=True;p.paragraph_format.page_break_before=(lev==1 and i!=63)
rm(old[132])
for i in range(30,61):rm(old[i])
P(29).style='TOC Heading';P(29).paragraph_format.page_break_before=True
after(old[29],'[[ASO_TOC]]')
for lev in [1,2,3]:
 s=d.styles[f'Heading {lev}'];s.font.name='Arial';s.font.size=Pt(13 if lev==1 else 11);s.font.bold=True;s.font.color.rgb=RGBColor(0,0,0)
 s.paragraph_format.space_before=Pt(12);s.paragraph_format.space_after=Pt(6)

# Copy the actual normative table from the ed4 template: grid, merged groups, shading, widths and labels.
ntxml=deepcopy(template.tables[0]._tbl);ts[0]._tbl.addnext(ntxml);rm(ts[0]._tbl);nt=Table(ntxml,d._body)
# Remove foreign bookmarks from the copied table to avoid duplicate bookmark identifiers.
for e in nt._tbl.xpath('.//w:bookmarkStart | .//w:bookmarkEnd'):rm(e)
for link in nt._tbl.xpath('.//w:hyperlink[@r:id]'):
 rel=template.part.rels[link.get(qn('r:id'))]
 link.set(qn('r:id'),d.part.relate_to(rel.target_ref,RT.HYPERLINK,is_external=True))
cell(nt.cell(5,3),'Orden 8/2025, de 22 de abril, de la Conselleria de Educación, Cultura, Universidades y Empleo, por la que se regula la evaluación del proceso de enseñanza-aprendizaje en ciclos formativos y cursos de especialización derivados de la Ley Orgánica 3/2022, de 31 de marzo, de ordenación e integración de la Formación Profesional. Modificada por la Orden 5/2026, de 1 de abril.',True)
cell(nt.cell(6,3),'Resolución de 16 de julio de 2026, de la Secretaría Autonómica de Educación, por la que se dictan instrucciones sobre ordenación académica y de organización de los centros que impartan los grados D y E de Formación Profesional durante el curso 2026-2027 en la Comunitat Valenciana.',True)
cell(nt.cell(7,3),'Resolución de 15 de junio de 2026, de la Dirección General de Centros Docentes, por la que se fija el calendario escolar del curso académico 2026-2027 en la Comunitat Valenciana, con las adaptaciones del calendario local aprobadas para Elche.',True)
cell(nt.cell(8,3),'Real Decreto 500/2024, de 21 de mayo, por el que se modifican determinados reales decretos por los que se establecen títulos de Formación Profesional de grado superior y se fijan sus enseñanzas mínimas, incluido el título de ASIR.',True)
cell(nt.cell(9,2),'Normativa de la modalidad presencial',True)
cell(nt.cell(9,3),'Resolución de 16 de julio de 2026, de instrucciones para los grados D y E de Formación Profesional durante 2026-2027, en sus disposiciones sobre organización y asistencia en régimen presencial. La Orden 30/2022 de semipresencial no se aplica a esta propuesta.',True)
cell(nt.cell(10,3),'Real Decreto 1629/2009, de 30 de octubre, por el que se establece el título de Técnico Superior en Administración de Sistemas Informáticos en Red y se fijan sus enseñanzas mínimas, con sus modificaciones.',True)
cell(nt.cell(11,3),'Decreto 114/2025, de 29 de julio, del Consell, por el que se establecen los currículos de los ciclos formativos de grado medio y superior. Tabla de ASIR, página 72: módulo 0374, 133 horas y 4 horas semanales. Orden 36/2012, de 22 de junio, como antecedente curricular y en lo que resulte aplicable conforme al marco vigente.',True)
x=after(nt._tbl,'Referencias complementarias: Real Decreto 69/2025, de 4 de febrero, sobre los elementos e instrumentos del Sistema Nacional de Formación Profesional; y Decreto 252/2019, de 29 de noviembre, sobre organización y funcionamiento de centros públicos de ESO, Bachillerato y Formación Profesional.')
for s in ['Fuentes: RD 500/2024 — https://www.boe.es/eli/es/rd/2024/05/21/500', 'Evaluación — https://dogv.gva.es/datos/2025/04/30/pdf/2025_13083_es.pdf', 'Instrucciones 2026/2027 — https://dogv.gva.es/datos/2026/07/20/pdf/2026_24495_es.pdf', 'Currículo y horario — https://dogv.gva.es/datos/2025/08/04/pdf/2025_29742_es.pdf']:
 x=after(x,s)

# Harmonize ECP nomenclature with the correspondence table already present.
lines=[]
for row in ts[4].rows[1:]:
 for line in row.cells[1].text.splitlines():
  line=line.strip()
  if line and line not in lines:lines.append(line)
cell(ts[1].cell(6,1),'Estándares de competencia del ciclo');cell(ts[1].cell(6,2),'\n'.join(lines));rm(ts[1].rows[7]._tr)
after(ts[1]._tbl,'Se utiliza la denominación de estándares de competencia profesional (ECP), coherente con el apartado 1.4. Las referencias UC de la normativa original se interpretan en su contexto histórico y no constituyen una segunda relación de competencias.')
cell(ts[5].cell(6,0),'Asociado a estándares de competencia');cell(ts[5].rows[6].cells[-1],'ECP0490_3: Gestionar servicios en el sistema informático.\nECP0485_3: Gestionar el «software» de un sistema informático.')

for row in ts[2].rows[1:-1]:
 name=row.cells[0].text.strip();lang='Inglés' if name.startswith('Inglés') else ('Castellano' if name.startswith('Administración de sistemas operativos') else 'Castellano (previsión)');cell(row.cells[1],lang)
after(ts[2]._tbl,'Lenguas vehiculares: se mantiene el castellano en ASO y el inglés en Inglés Profesional. Para el resto del ciclo se propone castellano como previsión de coordinación, que se adecuará al proyecto lingüístico y a la asignación del centro. El cuadro identifica un módulo propio de Inglés Profesional; no se atribuyen horas adicionales de inglés a ASO. Las actividades CLIL se integran en las horas previstas.')
put(90,'Duración curricular de referencia: 133 horas y 4 horas semanales en segundo curso, según el Decreto 114/2025, tabla de ASIR, página 72.')
put(91,'Horario de aula confirmado: miércoles y jueves, dos horas cada día. En el escenario del 09/09/2026 al 23/03/2027, la última sesión calculada es el 18/03/2027. Descontados los días no lectivos, se prevén 51 sesiones de dos horas: 102 horas, con 52 en la primera evaluación y 50 en la segunda. Las fechas de pruebas y evaluaciones se ajustarán al calendario del centro.')
put(92,'Se propone la siguiente distribución didáctica completa de las 133 horas: 92 horas de unidades en el aula, 10 de pruebas y recuperación y 31 de aplicación, consolidación y verificación de aprendizajes. La aplicación se desarrollará en la empresa cuando existan tareas adecuadas o mediante actividades supervisadas en el centro, conforme al plan individual y a la organización del equipo docente.')
units=[('UD1. Scripting',24,6,30),('UD2. Procesos',8,4,12),('UD3. Directorio',16,2,18),('UD4. Integración',16,8,24),('UD5. Automatización',16,5,21),('UD6. Administración remota',8,4,12),('UD7. Impresión',4,2,6),('Pruebas y recuperación',10,0,10),('TOTAL',102,31,133)]
th=table(old[92],['Bloque','Aula (h)','Aplicación (h)*','Total (h)'],units,[7.2,2.8,3.5,3.0])
after(th._tbl,'* Aplicación provisional en empresa o centro, según las oportunidades del destino y el plan formativo. Se priorizan integración, scripting y automatización; directorio e impresión pueden resolverse mediante prácticas de laboratorio. La distribución temporal no modifica los pesos de los RA.')
put(93,'Las 31 horas de aplicación son una propuesta didáctica para completar el módulo, no una duración de empresa fijada por la normativa ni una asignación automática por diferencia de calendario. El plan individual identificará su distribución entre centro y empresa y sus fechas. Si el destino no ofrece una tarea de ASO, el centro programará su alternativa con horario efectivo. No se sustituirán horas formativas por deberes sin supervisión ni se computarán dos veces las mismas horas.')
put(94,'La ventana general de formación en empresa comunicada por el centro va del 6 de abril al 18 de junio de 2027. Las 400 horas de empresa de segundo corresponden al conjunto del ciclo. La contribución concreta de ASO dependerá de las actividades acordadas y de las horas registradas; no se deduce de los porcentajes de calificación. Las fichas distinguen la base de aula y la fase de aplicación provisional.')
put(104,'El alumnado de Formación Profesional del centro presenta diversidad de edades, itinerarios académicos, experiencia laboral y competencia digital. En grado superior confluyen personas procedentes de Bachillerato, ciclos de grado medio, pruebas de acceso y otras vías previstas. La programación de aula concretará esta descripción general mediante la evaluación inicial, sin atribuir a todo el alumnado las características de un grupo particular.')
for i in range(105,111):rm(old[i])
after(old[104],'En ASO se comprobarán especialmente los conocimientos de Linux y Windows, terminal, redes, permisos y máquinas virtuales. Se ofrecerán guías graduadas, materiales accesibles y actividades de refuerzo y profundización. La tutoría y orientación coordinarán los apoyos de acuerdo con el PEC y con las necesidades detectadas.')
put(113,'El centro se sitúa en Elche, en un entorno industrial y de servicios donde conviven empresas del calzado, pymes, departamentos de informática y proveedores tecnológicos. La propuesta aproxima el aprendizaje a la administración de sistemas, al soporte a usuarios y a la continuidad de los servicios de estas organizaciones.')
put(154,'La propuesta concreta las líneas pedagógicas de este apartado mediante las siguientes actuaciones de ASO. Su coordinación con el apartado 4.1 del PEC y con el proyecto lingüístico se realizará en el departamento, conforme a los acuerdos del centro.')
for i,s in zip(range(155,165),[
'Inclusión y diversidad: guías graduadas, demostraciones, materiales accesibles, apoyos de orientación y distintas formas de presentar evidencias de los mismos CE.',
'Continuidad del aprendizaje: repaso inicial de Linux y Windows y reutilización progresiva del laboratorio entre scripting, directorio, integración, automatización y acceso remoto.',
'Aprendizaje competencial: prácticas con una necesidad de administración, verificación del resultado y documentación de decisiones e incidencias vinculadas a RA y CE.',
'Trabajo colaborativo: resolución de incidencias por parejas, rotación de funciones, revisión entre iguales y comprobación individual de los aprendizajes.',
'Proyecto intermodular: servidor y cliente de pruebas comunes y evidencias compartidas con el proyecto Proxmox, manteniendo la evaluación propia de ASO.',
'Uso de TIC: máquinas virtuales, Aules, documentación técnica, control de versiones y registros de pruebas para un trabajo reproducible.',
'Comunicación e idiomas: lectura guiada de documentación en inglés, glosarios y breves explicaciones técnicas, con apoyos lingüísticos.',
'Igualdad y orientación profesional: ejemplos diversos de profesionales, reparto equitativo de tareas y análisis de roles de administración y soporte.',
'Autonomía y actualización: búsqueda contrastada de documentación, justificación de soluciones y análisis de errores antes de aplicar cambios.',
'Sostenibilidad y responsabilidad: reutilización del hardware, dimensionamiento de máquinas virtuales, copias de seguridad y uso responsable de datos y credenciales.']):put(i,s)
put(168,'Recursos comunes: aula con equipos conectados en red, laboratorio de máquinas virtuales Linux y Windows, terminal y herramientas de administración gráfica, Aules para guías, entregas y rúbricas, y repositorio de documentación y scripts. Se emplearán las imágenes y licencias disponibles en el centro y redes de prácticas aisladas cuando sea necesario.')
after(old[168],'Documentación de consulta: manuales del sistema (man/help) y documentación oficial de GNU Bash, PowerShell, Windows Server, OpenLDAP, Samba, systemd, OpenSSH y CUPS. El profesorado publicará en Aules los enlaces y versiones utilizados. Para el proyecto se reutilizarán el entorno Proxmox y los materiales coordinados con PI.')

clil=[
'Lectura guiada de un ejemplo de Bash o PowerShell en inglés; glosario de variables, bucles, funciones y errores y comentario breve de un script.',
'Interpretación de ayuda y mensajes sobre procesos y servicios en inglés; glosario de process, thread, state, priority y service y explicación de una incidencia.',
'Lectura guiada de documentación LDAP o Active Directory en inglés; identificación de entry, attribute, group y filter y explicación de una consulta.',
'Consulta de documentación Samba en inglés; glosario de share, permission y access y descripción de una prueba cliente-servidor.',
'Lectura de ejemplos y mensajes de cron, systemd timers o Task Scheduler en inglés; interpretación de schedule, trigger y log y documentación de una tarea.',
'Consulta de documentación OpenSSH o acceso remoto en inglés; glosario de key, host, session y encryption y explicación de una conexión segura.',
'Lectura guiada de documentación CUPS y mensajes de cola en inglés; glosario de queue, job, printer y driver y descripción de una prueba de impresión.']
resources=[
'Linux, Windows y PowerShell; editor, terminal, manuales Bash/PowerShell y guías y rúbrica de scripting en Aules.',
'Máquinas virtuales Linux/Windows; herramientas de procesos, servicios, arranque y registros; casos de incidencia y guías en Aules.',
'Servidor y cliente de laboratorio; OpenLDAP y/o Active Directory, herramientas de consulta y administración; guías y evidencias en Aules.',
'Clientes Linux/Windows, Samba y red de prácticas; usuarios, grupos y recursos compartidos; matriz de pruebas de permisos en Aules.',
'Laboratorio reutilizado; cron/at, temporizadores systemd y programador de tareas de Windows; scripts de partida y registros.',
'Servidor y cliente de prácticas; OpenSSH y herramientas de escritorio remoto disponibles; claves de prueba y guía de verificación.',
'CUPS y servicio de impresión de Windows cuando esté disponible; impresora física o virtual, clientes y guía de gestión de colas.']
letters={1:'a',12:'l',13:'m',14:'n',15:'ñ',16:'o',18:'q',19:'r',20:'s'}
for u,ti in enumerate(range(12,19)):
 t=ts[ti]
 for ri in [5,7]:
  for ce in t.rows[ri]._tr.findall(qn('w:tc')):
   c=_Cell(ce,t);val=c.text
   val=re.sub(r'(?<!\d)(1|12|13|14|15|16|18|19|20)[a-zñ]?\.\s*',lambda m:('\n' if m.start()>0 else '')+letters[int(m.group(1))]+') ',val)
   val=re.sub(r'(?<!\d)2b?\.\s*','\nb) ',val);cell(c,val.strip())
 label=deepcopy(t.rows[17]._tr);content=deepcopy(t.rows[18]._tr)
 t.rows[19]._tr.addprevious(label);label.addnext(content)
 cell(_Cell(label.find(qn('w:tc')),t),'Actividades CLIL')
 cell(_Cell(content.find(qn('w:tc')),t),clil[u]+' Se integra en las horas de la unidad con apoyos lingüísticos y sin una ponderación de idioma independiente de los CE técnicos.')
 cell(t.rows[-1].cells[0],resources[u])
 base,app,total=units[u][1:];real=t.rows[0]._tr.findall(qn('w:tc'))
 cell(_Cell(real[-2],t),'Sesiones / horas aula');cell(_Cell(real[-1],t),f'{base//2} / {base} h')
 detail=deepcopy(t.rows[-1]._tr);t.rows[0]._tr.addnext(detail)
 cell(_Cell(detail.find(qn('w:tc')),t),f'Temporalización provisional: {base} h de aula ({base//2} sesiones de 2 h) + {app} h de aplicación en empresa o centro = {total} h. Véanse los apartados 1.3 y 6.')
 for p in t.cell(0,0).paragraphs:p.paragraph_format.page_break_before=False
before(ts[12]._tbl,'Las horas de aplicación se desarrollarán con las actividades del apartado 6 o su equivalente supervisado en el centro. El registro de aula y el plan individual identificarán ámbito, fechas y duración efectiva, sin duplicar horas.')
put(201,'La coordinación se establece con el módulo de Proyecto intermodular (PI), mediante la construcción, instalación y administración de un nodo Proxmox individual con máquinas virtuales o contenedores y servicios Linux automatizados. Cada estudiante trabaja con hardware reutilizado autorizado o con Proxmox virtualizado en VirtualBox. Se utiliza un entorno de prácticas independiente de la infraestructura de producción del centro.')

put(210,'PLAN PROVISIONAL Y ADAPTABLE A CADA EMPRESA. ASO participará mediante tareas de administración, soporte y documentación acordes con la actividad del destino. La tabla siguiente propone una selección; no obliga a que todas las empresas ofrezcan los siete RA. Antes de la incorporación, las tutorías del centro y de la empresa y el profesorado de ASO concretarán los RA y CE viables, tareas, entorno, fechas, horas, evidencias y personas responsables en el plan individual.')
put(211,'La ventana general de estancia comunicada por el centro va del 6 de abril al 18 de junio de 2027; las fechas y jornadas individuales se acordarán con cada empresa. Las 400 horas de segundo corresponden al conjunto del ciclo. La aplicación de ASO del apartado 1.3 se incorporará cuando existan tareas adecuadas; el resto se programará en el centro.')
put(212,'Selección según el destino: en empresas de soporte se priorizarán procesos, acceso remoto y documentación; en empresas con servidores se podrán incorporar recursos compartidos, automatización y scripts. Directorio e impresión se incluirán si forman parte de la actividad real y permiten intervenir con supervisión. Se admiten tecnologías equivalentes que acrediten el mismo CE. Los criterios que exigen sistemas heterogéneos o guiones en sistemas libres y propietarios se completarán en el centro cuando sea necesario.')
put(213,'Modo de evaluación compartida: los porcentajes propuestos representan el peso de los CE seleccionados dentro de cada RA, tomado del anexo I; no son porcentajes de horas ni una nota numérica emitida por la empresa. El centro trabaja la base de todos los RA y evalúa los CE restantes. Para los CE seleccionados, el docente integra las evidencias, la valoración de la tutoría de empresa y, si se necesita, una comprobación individual. Cada CE recibe una única calificación; no se duplican las notas de las prácticas de aula y de empresa.')
put(214,'Adaptación: si un RA no se activa en empresa, su reparto pasa a 100 % centro, 0 % empresa y 0 % ampliación. Si solo se activa parte de los CE propuestos, el porcentaje de empresa será la suma de sus pesos en el anexo I y el de centro, el complemento a 100 %. El plan y sus cambios se documentarán y comunicarán al alumnado antes de su aplicación. La falta de oportunidades en el destino no penaliza: se ofrece una alternativa de centro con el mismo CE y peso. Esta alternativa para ASO no sustituye la fase de empresa exigida al conjunto del ciclo.')
prs=[('RA1. Directorio','Sí','Condicional','RA1.d, i, j','70,00 %','30,00 %','0 %'),('RA2. Procesos','Sí','Propuesto','RA2.f, i','77,77 %','22,23 %','0 %'),('RA3. Automatización','Sí','Propuesto','RA3.d, h','75,00 %','25,00 %','0 %'),('RA4. Acceso remoto','Sí','Propuesto','RA4.e, i','77,77 %','22,23 %','0 %'),('RA5. Impresión','Sí','Condicional','RA5.g, i','77,77 %','22,23 %','0 %'),('RA6. Integración','Sí','Propuesto','RA6.b, f, h','62,50 %','37,50 %','0 %'),('RA7. Scripting','Sí','Propuesto','RA7.e, i','77,77 %','22,23 %','0 %')]
pt=table(ts[20]._tbl,['Resultado de aprendizaje','Centro','Empresa*','Modo: CE compartidos','% Centro','% Empresa','% Ampliac.'],prs,[3.3,1.5,2.5,3.4,1.95,1.95,1.9]);rm(ts[20]._tbl)
for c in pt.rows[0].cells:
 for p0 in c.paragraphs:
  for r0 in p0.runs:r0.font.size=Pt(9)
x=after(pt._tbl,'* Propuesto: se prioriza si el destino ofrece una tarea adecuada. Condicional: requiere directorio o impresión disponible. Ambos casos se concretan en el plan individual. El 0 % de ampliación indica que no se añaden bonificaciones ajenas a los CE; las tareas avanzadas se valoran, si procede, dentro de los mismos criterios.')
x=after(x,'Actividades formativas a realizar en la empresa relacionadas con los resultados de aprendizaje','Heading 3')
acts=[('RA1.d, i, j','Modificar con supervisión una cuenta, grupo o ajuste autorizado del directorio y documentar el cambio.','Registro anonimizado y comprobación del resultado.','Directorio de laboratorio, usuario y cliente de pruebas.',2),('RA2.f, i','Consultar procesos, consumo y registros; identificar un servicio y documentar su estado y función ante una incidencia.','Ficha de incidencia, herramientas y observaciones validadas.','Diagnóstico y seguimiento de procesos en máquina virtual.',4),('RA3.d, h','Configurar o revisar una tarea periódica de mantenimiento o copia y documentar su ejecución.','Planificación, registro y explicación de la comprobación.','Tarea programada en Linux o Windows con registros.',5),('RA4.e, i','Gestionar una intervención remota autorizada, revisar el servicio y documentar el procedimiento.','Parte de intervención sin credenciales y validación de la tutoría.','Acceso remoto al servidor de laboratorio y documentación.',4),('RA5.g, i','Gestionar una cola o trabajo de impresión y documentar configuración o resolución de una incidencia.','Registro de trabajo y ficha de configuración sin datos de usuarios.','CUPS o impresora virtual, cliente y trabajo de prueba.',2),('RA6.b, f, h','Revisar permisos de un recurso compartido, probar accesos autorizados y denegados y documentar la configuración.','Matriz de permisos y pruebas con datos ficticios o anonimizados.','Recurso Linux/Windows y pruebas de acceso.',8),('RA7.e, i','Crear o adaptar un guion de administración de un servicio, probarlo y documentar su uso.','Script sin datos sensibles, prueba controlada y explicación individual.','Guion de administración en el laboratorio y documentación.',6)]
at=table(x,['RA / CE','Actividad adaptable al destino','Evidencia prevista','Alternativa en el centro','Horas'],acts,[2.1,4.6,4.1,4.3,1.4]);x=at._tbl
for s in [
'Las horas orientativas suman 31 y corresponden a la fase de aplicación provisional. No se añaden a las 400 horas de empresa del ciclo ni son un reparto obligatorio por destino. La preparación inicial se realiza en las horas de aula de cada unidad. El plan individual distribuirá la aplicación entre empresa y centro y ajustará duraciones sin superar ni duplicar las 133 horas del módulo.',
'Seguimiento: reunión inicial para seleccionar tareas y medios; registro semanal breve de actividad del alumnado; contacto entre tutorías al menos quincenal o con mayor frecuencia ante incidencias; revisión intermedia y valoración final. La tutoría del centro y el docente de ASO registrarán el seguimiento en los instrumentos del centro y en ITACA SAO-Gestión cuando corresponda.',
'Evaluación: la tutoría de empresa valorará los RA previstos como superado o no superado y aportará observaciones cualitativas. El docente de ASO determinará la calificación numérica según los CE, las evidencias y las comprobaciones necesarias; «superado» no se transforma automáticamente en una nota fija. Se contrastarán cumplimiento del CE, funcionamiento verificado y documentación técnica. Los aprendizajes no superados se incorporarán a refuerzo o recuperación en el centro.',
'Información y evidencias: se utilizarán descripciones, datos de prueba y registros anonimizados autorizados. No se entregarán credenciales, datos de clientes ni configuraciones sensibles. Si no se pueden extraer evidencias, la tutoría de empresa podrá acreditar la tarea mediante informe y el docente podrá realizar una entrevista técnica o demostración equivalente.',
'Concreción individual: el plan identificará empresa y puesto o departamento, tutorías, RA/CE seleccionados, actividades y tecnologías, calendario y horas, evidencias autorizadas, porcentajes resultantes y alternativa de centro. Se revisará ante cambios de funciones o medios, manteniendo la trazabilidad y la evaluación objetiva.'
]:x=after(x,s)
put(218,'Evaluación inicial o diagnóstica. Se realiza al comienzo del módulo, en septiembre, para ajustar la planificación a las capacidades y necesidades del alumnado. Combina un cuestionario de conocimientos previos y tareas breves de terminal, permisos, redes y máquinas virtuales.')
put(238,'Actividades de recuperación complementarias que sustituyan a las no superadas, especialmente cuando ya se hayan facilitado soluciones, manteniendo los CE y el nivel de exigencia de la actividad original.').style='List Paragraph'
after(old[232],'Los CE desarrollados en empresa se valorarán según el apartado 6, manteniendo los pesos del anexo I. Si una actividad no se puede realizar por causas ajenas al alumnado, el mismo CE se evaluará mediante la alternativa de centro. No se añadirá un porcentaje de empresa por encima del 100 % del módulo.')
after(old[246],'El documento F15.PC02 será visado por Jefatura de estudios de Formación Profesional. La pérdida de evaluación continua no impide la evaluación final. Se favorecerá la reincorporación al seguimiento continuo cuando proceda conforme a la normativa y al procedimiento del centro, dejando constancia de las medidas y evidencias requeridas.')
put(273,'Nivel IV - Itinerario formativo personalizado en Formación Profesional: las medidas y, cuando proceda, la ampliación de convocatorias se tramitarán conforme a la normativa de evaluación y al asesoramiento de orientación; no se aplicarán de forma automática.')
x=after(ts[-1]._tbl,'Ponderación de los criterios de evaluación distribuidos entre diferentes unidades','Heading 3')
rt=table(x,['Criterios de evaluación','Unidad de referencia','Ponderación de cada CE en su unidad'],[('RA7.a–i','UD1','100 %'),('RA2.a–i','UD2','100 %'),('RA1.a–j','UD3','100 %'),('RA6.a–h','UD4','100 %'),('RA3.a–h','UD5','100 %'),('RA4.a–i','UD6','100 %'),('RA5.a–i','UD7','100 %')],[6.3,4,6.2])
after(rt._tbl,'Cada CE se atribuye al 100 % a una única unidad de referencia; no se reparte actualmente entre varias unidades. Sus pesos dentro del RA son los de las tablas anteriores. Las evidencias de PI y de empresa pueden contribuir a esos mismos CE y se integran en una única calificación. El reparto centro/empresa del apartado 6 identifica el origen de las evidencias y no modifica la ponderación curricular del anexo.')
fixes={'schellscript':'shell script','wondows':'Windows','powershell':'PowerShell','Cups':'CUPS','NextCloud':'Nextcloud','Secundaría':'Secundaria','Sistemas Informáticos en Redes':'Sistemas Informáticos en Red','personal especializados':'personal especializado','perspectivas de genero':'perspectivas de género','del módulo la unidad formativa':'del módulo','soluciones ,':'soluciones,'}
def fix(p):
 for r in p.runs:
  for a,b in fixes.items():r.text=r.text.replace(a,b)
for p in d.paragraphs:fix(p)
for t in d.tables:
 seen=set()
 for row in t.rows:
  for c in row.cells:
   if c._tc in seen:continue
   seen.add(c._tc)
   for p in c.paragraphs:
    # Keep the original normative table's alignment and typography.
    if t._tbl is not nt._tbl:p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    fix(p)
 pr=t.rows[0]._tr.get_or_add_trPr()
 if pr.find(qn('w:tblHeader')) is None:pr.append(OxmlElement('w:tblHeader'))
for el in list(body):
 if el.tag!=qn('w:p') or list(body).index(el)<list(body).index(old[61]):continue
 if not ''.join(el.itertext()).strip() and not el.xpath('.//w:sectPr | .//w:br | .//w:drawing | .//w:fldChar'):rm(el)
for p in d.paragraphs:
 if p.style.name.startswith('Heading'):
  nonum(p)
  if not p.text.strip():p.style='Normal'
for ti in range(12,19):
 e=before(ts[ti]._tbl,'');p=Paragraph(e,d._body);p.paragraph_format.page_break_before=True;p.paragraph_format.keep_with_next=True
 p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(1)
for a,b in zip(list(body),list(body)[1:]):
 if a.tag==qn('w:tbl') and b.tag==qn('w:tbl'):
  e=after(a,'');p=Paragraph(e,d._body);p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(2)
for row in ts[9].rows:
 for c in row.cells:
  for p in c.paragraphs:
   p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(2)
   for r in p.runs:r.font.size=Pt(10)
uf=d.settings.element.find(qn('w:updateFields'))
if uf is None:uf=OxmlElement('w:updateFields');d.settings.element.append(uf)
uf.set(qn('w:val'),'true')
# Finish filled fields in black, preserving white-on-grey table labels and blue hyperlinks.
for row in nt.rows[1:]:
 for c in row.cells[1:]:
  for p in c.paragraphs:
   for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
for p in d.paragraphs:
 for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
 if p.style.name.startswith('Heading'):
  pr=p._p.get_or_add_pPr()
  for e in list(pr):
   if e.tag==qn('w:pBdr'):pr.remove(e)
for lev in [1,2,3]:
 pr=d.styles[f'Heading {lev}']._element.pPr
 if pr is not None:
  for e in list(pr):
   if e.tag==qn('w:pBdr'):pr.remove(e)
d.core_properties.title='Propuesta didáctica · Administración de sistemas operativos · 2026/2027'
d.core_properties.subject='F01.PC02 edición 04. Plan de empresa provisional y adaptable.'
d.core_properties.comments='Revisión de 30/09/2026. Tabla normativa basada directamente en la plantilla ed4.'
d.save(W/'adaptada.docx')
print('Saved',W/'adaptada.docx','tables',len(d.tables))
